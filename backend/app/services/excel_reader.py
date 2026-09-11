from pathlib import Path
import openpyxl
from pyxlsb import open_workbook

class WorkbookReader:
    def __init__(self, file_like, filename):
        self.is_xlsb = Path(filename).suffix.lower() == ".xlsb"; file_like.seek(0)
        self.book = open_workbook(file_like) if self.is_xlsb else openpyxl.load_workbook(file_like, data_only=True, read_only=True)
        self.sheet_names = self.book.sheets if self.is_xlsb else self.book.sheetnames
    def require_sheet(self, sheet_name):
        value = next((name for name in self.sheet_names if name.strip().lower() == sheet_name.lower()), None)
        if not value: raise KeyError(sheet_name)
        return value
    def cell(self, sheet_name, row, col):
        return next((values.get(col) for _, values in self.iter_rows(sheet_name, row, row)), None)
    def iter_rows(self, sheet_name, min_row=1, max_row=None):
        name = self.require_sheet(sheet_name)
        if self.is_xlsb:
            with self.book.get_sheet(name) as sheet:
                for index, row in enumerate(sheet.rows(), 1):
                    if index >= min_row and (max_row is None or index <= max_row): yield index, {col + 1: cell.v for col, cell in enumerate(row) if cell.v is not None}
        else:
            for index, row in enumerate(self.book[name].iter_rows(min_row=min_row, max_row=max_row, values_only=True), min_row): yield index, {col + 1: value for col, value in enumerate(row) if value is not None}
    def close(self): self.book.close()
